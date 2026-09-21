"""
Generate Microsoft Word (.docx) Document for Patent Draft Application
Following the exact standard format from the patent drafting template,
incorporating the official patent CAD drawings (Figs 1-9) and updated reference numerals.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_patent_docx():
    doc = docx.Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base style
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0, 0, 0)

    # Document Header Title
    header_p = doc.add_paragraph()
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = header_p.add_run("PATENT APPLICATION DRAFT")
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = 'Times New Roman'

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("Following information is required for drafting of patent application.")
    r_sub.italic = True
    r_sub.font.size = Pt(11)
    
    doc.add_paragraph()

    def add_section_heading(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r_num = p.add_run(num_str + " ")
        r_num.bold = True
        r_num.font.size = Pt(12)
        r_title = p.add_run(title_str)
        r_title.bold = True
        r_title.font.size = Pt(12)
        return p

    # -------------------------------------------------------------
    # 1. Full name, nationality and address of applicant(s):
    # -------------------------------------------------------------
    add_section_heading("1.", "Full name, nationality and address of applicant(s):")
    
    table1 = doc.add_table(rows=3, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    col_widths = [Inches(2.5), Inches(1.2), Inches(2.8)]
    headers1 = ["Full Name", "Nationality", "Address"]
    
    hdr_cells = table1.rows[0].cells
    for i, name in enumerate(headers1):
        hdr_cells[i].text = name
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr_cells[i], 'E0E0E0')

    data1 = [
        ("Vishwakarma Institute of Technology", "Indian", "666, Upper Indiranagar, Bibwewadi, Pune, Maharashtra, India – 411 037"),
        ("Vishwakarma University", "Indian", "Survey No 2, 3, 4, Kondhwa Main Rd, Laxmi Nagar, Betal Nagar, Kondhwa, Pune, Maharashtra, India - 411048")
    ]
    for row_idx, row_data in enumerate(data1):
        row_cells = table1.rows[row_idx + 1].cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text

    for row in table1.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_border(cell, 
                            top=dict(sz=4, val='single', color='000000'),
                            bottom=dict(sz=4, val='single', color='000000'),
                            left=dict(sz=4, val='single', color='000000'),
                            right=dict(sz=4, val='single', color='000000'))

    doc.add_paragraph()

    # -------------------------------------------------------------
    # 2. Full name, nationality, address, mail id, and phone number of inventor(s):
    # -------------------------------------------------------------
    add_section_heading("2.", "Full name (including middle name), nationality, address, mail id, and phone number of inventor(s):")

    table2 = doc.add_table(rows=3, cols=5)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False
    
    col_widths2 = [Inches(1.5), Inches(0.6), Inches(2.2), Inches(1.4), Inches(0.8)]
    headers2 = ["Full Name (Including middle name)", "Nationality", "VIT Address (Start with full dept. name followed by full institute name)", "Mail ID", "Phone No."]

    hdr_cells2 = table2.rows[0].cells
    for i, name in enumerate(headers2):
        hdr_cells2[i].text = name
        hdr_cells2[i].paragraphs[0].runs[0].bold = True
        hdr_cells2[i].paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(hdr_cells2[i], 'E0E0E0')

    data2 = [
        ("Prof. Vijaykumar Raghunath Ghule", "IN", "Department of Computer Science & Engineering (Data Science), Vishwakarma Institute of Technology, Pune", "Vijaykumar.ghule@vit.edu", "8624046893"),
        ("Sail Sitaram Nagale", "IN", "Department of Computer Science & Engineering (Data Science), Vishwakarma Institute of Technology, Pune", "sail.nagale@vit.edu", "")
    ]
    for row_idx, row_data in enumerate(data2):
        row_cells = table2.rows[row_idx + 1].cells
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = text
            row_cells[col_idx].paragraphs[0].runs[0].font.size = Pt(9.5)

    for row in table2.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths2[i]
            set_cell_border(cell, 
                            top=dict(sz=4, val='single', color='000000'),
                            bottom=dict(sz=4, val='single', color='000000'),
                            left=dict(sz=4, val='single', color='000000'),
                            right=dict(sz=4, val='single', color='000000'))

    doc.add_paragraph()

    # Signatures Table
    sig_p = doc.add_paragraph()
    sig_p.add_run("Copy and paste clear soft copy of signatures of all inventors in the following table: (Insert or delete the cells as per number of inventors)").italic = True
    
    table_sig = doc.add_table(rows=2, cols=4)
    table_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_sig.autofit = False
    
    sig_names = ["Vijaykumar Ghule", "Sail Sitaram Nagale", "Name 3", "Name 4"]
    for i, cell in enumerate(table_sig.rows[0].cells):
        cell.text = sig_names[i]
        cell.paragraphs[0].runs[0].bold = True
        set_cell_background(cell, 'F0F0F0')
    
    for cell in table_sig.rows[1].cells:
        cell.text = "\n\n(Signature)\n"

    for row in table_sig.rows:
        for cell in row.cells:
            cell.width = Inches(1.6)
            set_cell_border(cell, 
                            top=dict(sz=4, val='single', color='000000'),
                            bottom=dict(sz=4, val='single', color='000000'),
                            left=dict(sz=4, val='single', color='000000'),
                            right=dict(sz=4, val='single', color='000000'))

    doc.add_paragraph()

    # -------------------------------------------------------------
    # 3. Title of the invention:
    # -------------------------------------------------------------
    add_section_heading("3.", "Title of the invention:")
    p3 = doc.add_paragraph()
    p3.paragraph_format.left_indent = Inches(0.25)
    r = p3.add_run("IOT AND MACHINE LEARNING-BASED PREDICTIVE WORKER SAFETY HELMET WITH REAL-TIME HAZARD DETECTION")
    r.bold = True
    
    p3_note = doc.add_paragraph()
    p3_note.paragraph_format.left_indent = Inches(0.25)
    p3_note.add_run("(Word count: 14 words - Complies with limit of not more than 15 words)").italic = True

    # -------------------------------------------------------------
    # 4. Technical field of the invention:
    # -------------------------------------------------------------
    add_section_heading("4.", "Technical field of the invention:")
    p4 = doc.add_paragraph()
    p4.paragraph_format.left_indent = Inches(0.25)
    p4.add_run(
        "The present invention relates generally to industrial occupational safety systems, wearable embedded electronic devices, and edge computing. "
        "More specifically, the present invention relates to an Internet of Things (IoT) and Machine Learning (ML) enabled predictive smart helmet capable of "
        "multi-modal environmental hazard detection, worker physiological biometric monitoring, 6-axis kinematic fall and impact detection, "
        "on-device predictive risk scoring, and dual-mode resilient wireless telemetry dispatch for construction, mining, and hazardous industrial sites."
    )

    # -------------------------------------------------------------
    # 5. Prior art:
    # -------------------------------------------------------------
    add_section_heading("5.", "Prior art:")
    p5 = doc.add_paragraph()
    p5.paragraph_format.left_indent = Inches(0.25)
    p5.add_run(
        "Conventional safety headgear used in industrial, construction, and mining facilities consists predominantly of passive thermoplastic or fiberglass shells engineered solely to cushion mechanical impact from falling debris. "
        "However, these traditional helmets cannot sense invisible hazards such as dangerous ambient toxic gases, prolonged extreme thermal exposure, or acute physiological distress in workers.\n\n"
        "Recent electronic safety helmets integrate single-variable sensors but suffer from critical drawbacks:\n"
        "1. Static Thresholding & High False Alarms: Existing devices trigger alarms based on rigid, point-in-time threshold breaches, failing to detect gradual pre-incident hazard buildup or generating nuisance alarms.\n"
        "2. Lack of Predictive Multi-Sensor Fusion: Conventional systems do not correlate physiological metrics (heart rate) with ambient parameters (temperature, toxic gases) and kinematic vectors (acceleration magnitude, jerk).\n"
        "3. Cloud Dependency & Network Latency: Existing solutions rely on constant cloud connectivity for computational inference, rendering them ineffective in radio-shadowed environments like tunnels or deep basements.\n"
        "4. Absence of Integrated Geo-Telemetry: Inability to pinpoint exact coordinates of incapacitated workers during emergencies delays rescue intervention.\n\n"
        "The present invention overcomes all of the above disadvantages by executing on-device edge ML inference, dynamic feature extraction, predictive risk scoring (0-100), dual-action local/remote alarming, and hybrid LoRa/Wi-Fi telemetry with GPS tracking."
    )

    # -------------------------------------------------------------
    # 6. Objective(s) of Invention:
    # -------------------------------------------------------------
    add_section_heading("6.", "Objective(s) of Invention:")
    objectives = [
        "To provide a non-intrusive smart safety helmet that continuously monitors environmental hazards, worker biometrics, and kinematic motion in real time.",
        "To provide an edge computing architecture with an embedded machine learning model capable of computing a dynamic continuous Safety Risk Score (0–100) and classifying conditions into Low, Medium, and High Risk tiers.",
        "To reliably detect emergencies including sudden falls, high-g mechanical impacts, hazardous gas accumulation, elevated heat stress, and abnormal heart rates before fatal incidents occur.",
        "To implement a dual-mode low-latency emergency response mechanism comprising instant localized audiovisual alarms (<100 ms) and automated wireless telemetry dispatch to a supervisory dashboard.",
        "To incorporate geospatial coordinates via an onboard GPS receiver to enable immediate search-and-rescue localization across large-scale industrial infrastructures.",
        "To log multi-modal time-series safety records to a central database for retrospective incident analysis, compliance audits, and proactive workplace hazard mitigation."
    ]
    for obj in objectives:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.5)
        p.add_run(obj)

    # -------------------------------------------------------------
    # 7. Synopsis:
    # -------------------------------------------------------------
    add_section_heading("7.", "Synopsis:")
    p7 = doc.add_paragraph()
    p7.paragraph_format.left_indent = Inches(0.25)
    p7.add_run(
        "The present invention is an IoT and Machine Learning-enabled Predictive Worker Safety Helmet. "
        "The device comprises an industrial protective helmet shell (100) embedding a Global Positioning System (GPS) module (101) mounted at the top crown, "
        "a dual temperature and hazardous gas sensor unit (102) positioned on the forward lateral housing, an audiovisual warning unit (103, piezoelectric buzzer and high-intensity LED) positioned on the front visor brim, "
        "an ergonomic chin strap (104) with an embedded photoplethysmography (PPG) pulse sensor for continuous pulse rate acquisition, "
        "a 6-axis Inertial Measurement Unit (105, IMU containing 3-axis accelerometer and 3-axis gyroscope) mounted on the lateral shell, "
        "and a central microcontroller and edge processing unit (106, ESP32 SoC with Wi-Fi/LoRa transceiver and power management circuitry) housed in a protected rear/lateral enclosure.\n\n"
        "The microcontroller samples multi-sensor feeds, applies edge noise filtering, and extracts dynamic statistical features (acceleration magnitude, jerk derivative, rate of temperature change, and gas concentration). "
        "A trained machine learning model running on the edge processor computes a real-time Risk Score (0–100). When the score exceeds predefined safety bounds or a critical fall/impact signature is recognized, "
        "the helmet immediately actuates local audiovisual warnings (103) and dispatches an emergency telemetry packet containing worker ID, physiological vitals, and GPS coordinates (101) to a central supervisory dashboard."
    )

    # -------------------------------------------------------------
    # 8. Brief description of drawings:
    # -------------------------------------------------------------
    add_section_heading("8.", "Brief description of drawings (if any):")
    p8 = doc.add_paragraph()
    p8.paragraph_format.left_indent = Inches(0.25)
    p8.add_run(
        "The invention is illustrated in the accompanying technical drawings (Figs. 1 to 9), wherein single common numerals are used to designate specific structural and electronic components across all views:\n\n"
        "• Figure 1 illustrates a Perspective View of the assembled predictive worker safety helmet.\n"
        "• Figure 2 illustrates a Front View of the predictive worker safety helmet.\n"
        "• Figure 3 illustrates a Side View of the predictive worker safety helmet.\n"
        "• Figure 4 illustrates a Top View showing crown component placement.\n"
        "• Figure 5 illustrates a Rear View showing the main controller enclosure.\n"
        "• Figure 6 illustrates a Bottom View (Inside) displaying internal harness and strap arrangement.\n"
        "• Figure 7 illustrates a Sectional View (taken along section line A–A).\n"
        "• Figure 8 illustrates an Exploded View depicting the mechanical and electronic modular assembly.\n"
        "• Figure 9 illustrates Individual Component Views of constituent sub-assemblies.\n\n"
        "Reference Numerals Denoting Components in Drawings:\n"
        "100 denotes Helmet Shell (Protective outer casing)\n"
        "101 denotes GPS Module (Geospatial positioning unit)\n"
        "102 denotes Temperature & Gas Sensor (Multi-modal environmental sensing module)\n"
        "103 denotes Buzzer & LED (Audiovisual alarm & optical warning unit)\n"
        "104 denotes Chin Strap (Ergonomic harness with integrated biometric pulse sensor)\n"
        "105 denotes IMU (6-axis Accelerometer & Gyroscope kinematic sensor)\n"
        "106 denotes ESP32 Controller (Main edge processing & wireless telemetry unit)"
    )

    # -------------------------------------------------------------
    # 9. Detailed description of the invention:
    # -------------------------------------------------------------
    add_section_heading("9.", "Detailed description of the invention:")
    p9 = doc.add_paragraph()
    p9.paragraph_format.left_indent = Inches(0.25)
    p9.add_run(
        "Referring to Figures 1 through 9, the predictive worker safety helmet comprises an industrial high-impact thermoplastic outer shell (100) engineered to provide structural head protection. "
        "A GPS module (101) is rigidly affixed to the top crown apex to maintain an unobstructed line-of-sight satellite reception for real-time spatial positioning.\n\n"
        "Mounted on the front lateral perimeter is a combined environmental sensor suite (102) comprising a digital temperature sensor and a toxic gas sensor configured to measure ambient heat and dangerous concentrations of CO, CH4, H2S, and LPG. "
        "On the front brim/visor, an audiovisual warning unit (103) is integrated, comprising an acoustic piezoelectric buzzer and a high-luminescence flashing alert LED for immediate wearer warning.\n\n"
        "Secured to the helmet shell is an adjustable chin strap (104) having an integrated biometric photoplethysmography (PPG) pulse sensor arranged to maintain continuous contact with the wearer's skin at the jaw/neck to measure cardiovascular pulse rate. "
        "A 6-axis Inertial Measurement Unit (105, IMU combining a 3-axis accelerometer and a 3-axis gyroscope) is mounted on the lateral shell to continuously track kinematic motion, tilt angles, free-fall dynamics, and impact vectors.\n\n"
        "The main processing and communication hub is an ESP32 microcontroller unit (106) mounted within a sealed, impact-resistant rear housing. The controller integrates a dual-core 32-bit CPU, power regulation circuitry, rechargeable lithium battery, and a hybrid Wi-Fi / LoRa wireless transceiver.\n\n"
        "The operational sequence and edge machine learning workflow are executed as follows:\n"
        "1. Real-Time Multi-Sensor Acquisition: The IMU (105) is sampled at 50 Hz, biometric sensor on chin strap (104) at 20 Hz, and temperature/gas sensors (102) at 1 Hz.\n"
        "2. Dynamic Edge Preprocessing & Feature Extraction: The controller (106) calculates total acceleration magnitude Amag(t) = sqrt(ax^2 + ay^2 + az^2), dynamic jerk derivative J(t) = d(Amag)/dt, moving average heart rate, temperature rate-of-change, and toxic gas slope.\n"
        "3. Predictive Risk Scoring (0–100): An onboard pre-trained machine learning model processes the feature vector to generate a continuous Safety Risk Score (R). The classification logic is partitioned into:\n"
        "   - Low Risk (0 <= R < 40): Normal operating state. Sensor logs transmitted periodically.\n"
        "   - Medium Risk (40 <= R < 75): Advisory state (e.g. rising heat or moderate gas). Intermittent warning signaled via LED (103).\n"
        "   - High Risk (75 <= R <= 100): Critical emergency (e.g. fall trajectory characterized by high jerk J > 25 g/s followed by immobility, acute gas leak > 50 PPM). Continuous local alarm triggered.\n"
        "4. Dual-Action Alert & Telemetry Dispatch: The controller triggers continuous audible alarm via buzzer/LED (103) and broadcasts an emergency packet containing worker ID, telemetry, and GPS coordinates (101) over LoRa/Wi-Fi to the supervisory command center."
    )

    # -------------------------------------------------------------
    # 10. Best method of performance of the invention:
    # -------------------------------------------------------------
    add_section_heading("10.", "Best method of performance of the invention:")
    p10 = doc.add_paragraph()
    p10.paragraph_format.left_indent = Inches(0.25)
    p10.add_run(
        "In practical deployment, an industrial worker fastens the helmet shell (100) securely using chin strap (104) and powers on the controller unit (106). "
        "The controller executes an automatic power-on self-test diagnostic on all sensors (101, 102, 104, 105), emits a confirmation beep via buzzer (103), and establishes a wireless handshake with the local gateway.\n\n"
        "During work, the IMU (105) continuously monitors head dynamics at 50 Hz. If the worker slips from scaffolding, the IMU records free-fall weightlessness followed by an impact shock (>3.5g) and a motionless window. "
        "Simultaneously, if toxic gas levels exceed safety thresholds or thermal stress causes heart rate escalation (>130 bpm), the edge ML model computes a compound risk score of R >= 85 (High Risk). "
        "Within 80 milliseconds, the buzzer and LED (103) fire at 90 dB, and an emergency LoRa packet containing latitude and longitude from the GPS module (101) is broadcast to the supervisory command dashboard, "
        "enabling rescue teams to pinpoint and extract the worker within minutes."
    )

    # -------------------------------------------------------------
    # 11. CLAIMS:
    # -------------------------------------------------------------
    add_section_heading("11.", "CLAIMS:")
    p11_lead = doc.add_paragraph()
    p11_lead.paragraph_format.left_indent = Inches(0.25)
    p11_lead.add_run("We Claim:").bold = True

    claims = [
        "1. An IoT and machine learning-enabled predictive worker safety helmet system comprising:\n"
        "a protective helmet shell (100) configured to be worn by an industrial worker;\n"
        "a geospatial positioning unit comprising a global positioning system (GPS) module (101) mounted to said helmet shell;\n"
        "an environmental sensing unit including a temperature and hazardous gas sensor module (102) mounted to said helmet shell;\n"
        "an audiovisual warning unit (103) comprising an acoustic buzzer and a visual light emitting diode (LED) positioned on said helmet shell;\n"
        "a chin strap (104) having an integrated biometric photoplethysmography (PPG) pulse sensor positioned to maintain skin contact;\n"
        "a kinematic sensing unit comprising a 6-axis inertial measurement unit (IMU) (105) including a 3-axis accelerometer and a 3-axis gyroscope; and\n"
        "a microcontroller and edge processing unit (106) communicatively coupled to said sensors, warning unit, and an integrated wireless transceiver, wherein said edge processing unit is configured to extract multi-modal feature vectors from real-time sensor data, compute a predictive safety risk score (0 to 100) using an embedded machine learning model, categorize said score into a plurality of risk tiers, and autonomously actuate said audiovisual warning unit (103) and transmit emergency telemetry packets to a supervisory dashboard upon detection of a hazardous condition.",

        "2. The system as claimed in claim 1, wherein said machine learning model is trained on historical multi-modal industrial sensor data and evaluates an input feature vector comprising instantaneous acceleration magnitude, dynamic jerk derivative, moving average heart rate, rate of temperature increase, and toxic gas concentration to predict potential safety incidents before threshold failure.",

        "3. The system as claimed in claim 1, wherein said edge processing unit (106) is configured with a tri-level risk classification logic comprising a Low Risk level (0 <= R < 40) signifying normal operational parameters, a Medium Risk level (40 <= R < 75) triggering an advisory alert on said supervisory dashboard and intermittent local warning, and a High Risk level (75 <= R <= 100) triggering continuous local acoustic and visual alarms and instantaneous transmission of emergency geospatial and physiological payloads.",

        "4. The system as claimed in claim 1, wherein said kinematic sensing unit (105) executes a fall and collision detection routine configured to identify a free-fall acceleration signature followed by a high-g impact peak exceeding a defined acceleration threshold and a subsequent zero-motion window.",

        "5. The system as claimed in claim 1, wherein said hazardous gas sensor (102) is configured to detect one or more toxic and flammable gases selected from the group consisting of Carbon Monoxide (CO), Methane (CH4), Hydrogen Sulfide (H2S), and Liquefied Petroleum Gas (LPG).",

        "6. The system as claimed in claim 1, wherein said wireless transceiver of controller (106) is configured with automated failover logic to transmit emergency packets over LoRa frequencies when local Wi-Fi network connectivity is unavailable.",

        "7. The system as claimed in claim 1, wherein said emergency telemetry packet transmitted to said supervisory dashboard comprises worker identification code, real-time GPS coordinates from module (101), time stamp, current heart rate from strap sensor (104), ambient gas concentration and temperature from module (102), and current predictive risk score.",

        "8. A computer-implemented method for predictive occupational safety monitoring using a smart helmet, comprising the steps of:\n"
        "acquiring continuous time-synchronized data streams from a multi-sensor array comprising temperature and gas sensor (102), biometric sensor (104), 6-axis IMU (105), and GPS module (101);\n"
        "filtering raw sensor data to eliminate high-frequency noise and motion artifacts;\n"
        "extracting dynamic temporal and kinematic features from the filtered sensor data;\n"
        "executing an embedded machine learning predictive risk assessment algorithm on the extracted features to generate a normalized numerical risk score from 0 to 100;\n"
        "comparing said risk score against defined safety classification thresholds;\n"
        "energizing a local audiovisual alert unit (103) on the helmet when the risk score exceeds a predetermined threshold; and\n"
        "transmitting an emergency telemetry packet containing geospatial location from GPS (101) and physiological metrics over a wireless interface to a remote monitoring dashboard for emergency response coordination."
    ]

    for c in claims:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(6)
        p.add_run(c)

    # -------------------------------------------------------------
    # 12. Inventive step of your invention:
    # -------------------------------------------------------------
    add_section_heading("12.", "Inventive step of your invention:")
    p12 = doc.add_paragraph()
    p12.paragraph_format.left_indent = Inches(0.25)
    p12.add_run(
        "The inventive step of the present invention resides in the technical convergence of:\n"
        "1. Multi-Modal Predictive Sensor Fusion vs. Static Thresholding: Instead of relying on isolated point-in-time threshold triggers, the invention dynamically correlates multi-channel physiological, environmental, and kinematic variables through an edge-optimized ML model, detecting evolving pre-incident hazards (e.g., heat stress combined with toxic gas inhalation) before worker incapacitation.\n"
        "2. On-Device Edge Processing & Sub-Second Latency: Embedding the feature extraction and classification model directly on the microcontroller (106) guarantees instantaneous local warning actuation without relying on cloud availability or suffering network delays.\n"
        "3. Resilient Dual-Mode Wireless Telemetry & Integrated GPS: Combining Wi-Fi with sub-GHz LoRa modulation ensures reliable emergency data transmission across deep construction trenches, tunnels, and remote industrial complexes alongside precise GPS coordinates from module (101)."
    )

    # -------------------------------------------------------------
    # 13. Industrial application:
    # -------------------------------------------------------------
    add_section_heading("13.", "Industrial application:")
    p13 = doc.add_paragraph()
    p13.paragraph_format.left_indent = Inches(0.25)
    p13.add_run(
        "The present invention has direct industrial utility across multiple hazardous sectors:\n"
        "• Civil Construction and Infrastructure: Real-time detection of falls from scaffolds, crane zone collisions, and heat exhaustion.\n"
        "• Mining and Tunneling: Continuous monitoring of toxic and explosive gas buildup (CO, CH4, H2S), asphyxiation risks, and worker entrapment.\n"
        "• Petrochemical Refineries & Chemical Plants: Early warning for volatile chemical leaks, excessive thermal proximity, and worker distress.\n"
        "• Heavy Manufacturing, Foundries & Metallurgy: Heat stress indexing and impact monitoring in high-temperature blast furnace zones.\n"
        "• Firefighting and Disaster Rescue: Real-time physiological tracking and geospatial tracing of emergency personnel in obscured environments."
    )

    # -------------------------------------------------------------
    # 14. Abstract:
    # -------------------------------------------------------------
    add_section_heading("14.", "Abstract:")
    p14 = doc.add_paragraph()
    p14.paragraph_format.left_indent = Inches(0.25)
    p14.add_run(
        "An IoT and machine learning-enabled predictive worker safety helmet and method are disclosed for real-time hazard detection and proactive occupational safety monitoring. "
        "The smart helmet comprises an industrial helmet shell (100) integrating a top-mounted GPS module (101), a temperature and hazardous gas sensor (102), "
        "a front-mounted audiovisual alarm (103), a chin strap with integrated biometric pulse sensor (104), a lateral 6-axis IMU (105), and an ESP32 edge microcontroller (106) with dual-mode wireless communication. "
        "The edge microcontroller continuously samples multi-modal sensor feeds, removes artifacts, extracts dynamic features, and executes an onboard machine learning model to compute a real-time risk score (0–100) "
        "categorized into low, medium, and high risk levels. Upon detecting hazardous events—such as falls, toxic gas leaks, excessive thermal stress, or cardiac irregularities—the helmet instantly actuates local alarms (103) "
        "and dispatches emergency telemetry packets containing exact GPS coordinates (101) to a central supervisory dashboard for rapid rescue intervention."
    )

    # -------------------------------------------------------------
    # 15. Drawing:
    # -------------------------------------------------------------
    add_section_heading("15.", "Drawing:")
    
    p15_note = doc.add_paragraph()
    p15_note.paragraph_format.left_indent = Inches(0.25)
    p15_note.add_run("Technical patent drawings containing Figs. 1 to 9 with standard single-numeral callouts (100–106):").italic = True

    # Embed drawing image
    img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figures", "patent_drawings.jpg")
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(8)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.5))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Sheet 1/1: Figures 1 to 9 (Perspective, Orthographic, Sectional, Exploded, and Component Views with Reference Numerals)")
        r_cap.bold = True
        r_cap.font.size = Pt(10)

    # Save Word document
    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Patent_Draft_Predictive_Worker_Safety_Helmet.docx")
    doc.save(output_path)
    print(f"Successfully generated Word (.docx) document at: {output_path}")

if __name__ == '__main__':
    create_patent_docx()
